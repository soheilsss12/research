#!/usr/bin/env python3
"""Mark P-10..P-19 as independently documented in the 41×43 control matrix.

This does not turn missing customer/price/market data into completed research. Each of those
rows remains explicitly marked as data-gated in the dossier text and in column I.
"""
from pathlib import Path
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
PATH = ROOT / "reports" / "ماتریس_کنترل_پوشش_۴۱_گزارش_مد_و_پوشاک.xlsx"
SHEET = "ماتریس پوشش ۴۱ پرونده"
PERSIAN = str.maketrans("0123456789", "۰۱۲۳۴۵۶۷۸۹")
IDS = {("پ-" + f"{n:02d}").translate(PERSIAN): n for n in range(10, 20)}

wb = load_workbook(PATH)
ws = wb[SHEET]
changed = 0
for row in range(5, ws.max_row + 1):
    pid = ws.cell(row, 1).value
    if pid not in IDS:
        continue
    n = IDS[pid]
    ws.cell(row, 7).value = "پروندهٔ مستقل اسنادی ۰٫۱؛ دادهٔ مفقود صریح/گیت ثبت شده"
    ws.cell(row, 8).value = f"dossier {pid} و Word اصلی؛ دفتر منبع P{n:02d}-S.."
    if n == 18:
        remaining = "خارج از دامنهٔ فعلی؛ فقط reversal صریح برای عملیات/لجستیک، سپس ارزیابی تازه"
    else:
        remaining = "پژوهش اولیه واقعی: حساب هدف، مشاهدهٔ جریان کار، شواهد قیمت/تمایل به پرداخت و ورودی‌های TAM/SAM/SOM"
    ws.cell(row, 9).value = remaining
    changed += 1
wb.save(PATH)
print(f"updated {changed} control rows in {PATH}")
